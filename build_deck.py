# -*- coding: utf-8 -*-
"""Build the VOIS major-project submission deck from the official template."""
import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = Path('data/VOIS_Major_Project_PPT_Submission_Template.pptx')
OUT      = Path('VOIS_Major_Project_Seasonal_Agriculture_Analysis.pptx')
FIG      = Path('figures')

FONT   = 'Trebuchet MS'
INK    = RGBColor(0x1B, 0x1B, 0x1A)
INK2   = RGBColor(0x45, 0x44, 0x41)
MUTED  = RGBColor(0x6F, 0x6E, 0x6A)
TEAL   = RGBColor(0x1B, 0xAF, 0x7A)
BLUE   = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
RED    = RGBColor(0xE3, 0x49, 0x48)

SW, SH = Inches(13.333), Inches(7.5)

# --------------------------------------------------------------------------- helpers
RT_SLIDE_LAYOUT = ('http://schemas.openxmlformats.org/officeDocument/2006/'
                   'relationships/slideLayout')
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def duplicate_slide(prs, idx):
    """Deep-copy a slide (shapes + relationships) and append it to the deck.

    python-pptx assigns its own rIds, so the copied shape XML is rewritten to point
    at the new ones - otherwise the duplicated slide loses its images.
    """
    src = prs.slides[idx]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):                      # drop layout-injected placeholders
        shp._element.getparent().remove(shp._element)

    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype == RT_SLIDE_LAYOUT:            # add_slide already wired this up
            continue
        target = rel.target_ref if rel.is_external else rel._target
        rid_map[rid] = dst.part.rels._add_relationship(rel.reltype, target, rel.is_external)

    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for name, value in list(node.attrib.items()):
                if name.startswith(f'{{{R_NS}}}') and value in rid_map:
                    node.set(name, rid_map[value])
        dst.shapes._spTree.append(el)
    return dst


def reorder(prs, order):
    """Reorder slides to the given list of current indices."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for e in ids:
        lst.remove(e)
    for i in order:
        lst.append(ids[i])


def shape_by_text(slide, needle):
    for shp in slide.shapes:
        if shp.has_text_frame and needle.lower() in shp.text_frame.text.lower():
            return shp
    return None


def set_text(shape, blocks, align=PP_ALIGN.LEFT, anchor=None):
    """blocks = [(text, size, bold, color, space_before, level), ...]"""
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.clear()
    for i, b in enumerate(blocks):
        text, size, bold, color = b[0], b[1], b[2], b[3]
        before = b[4] if len(b) > 4 else 0
        level  = b[5] if len(b) > 5 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = level
        if before:
            p.space_before = Pt(before)
        r = p.add_run(); r.text = text
        f = r.font
        f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    return shape


def add_box(slide, l, t, w, h, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    set_text(box, blocks, align=align, anchor=anchor)
    return box


def add_fitted_picture(slide, path, box_l, box_t, box_w, box_h):
    """Insert an image scaled to fit the box, centred."""
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    w = box_w
    h = w / ar
    if h > box_h:
        h = box_h
        w = h * ar
    left = box_l + (box_w - w) / 2
    top  = box_t + (box_h - h) / 2
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                    Inches(w), Inches(h))


def clear(shape):
    tf = shape.text_frame
    tf.clear()
    tf.paragraphs[0].runs and None
    return shape


def results_slide(slide, title, headline, image, takeaway, img_top=1.42, img_h=4.95,
                  takeaway_on_top=False):
    """Lay out one RESULTS slide: title, headline, figure, takeaway strip."""
    # title stays "RESULTS"; widen it so a subtitle fits beside it
    t = shape_by_text(slide, 'RESULTS')
    if t is not None:
        t.width = Inches(4.2)
        set_text(t, [('RESULTS', 30, True, INK)])
    # remove the template's leftover prompt boxes
    for shp in list(slide.shapes):
        if shp.has_text_frame and 'screen shots' in shp.text_frame.text.lower():
            shp._element.getparent().remove(shp._element)
    for shp in list(slide.shapes):
        if (shp.has_text_frame and shp.text_frame.text.strip() == ''
                and shp.name.startswith('Text Placeholder 30')):
            shp._element.getparent().remove(shp._element)

    if takeaway_on_top:
        add_box(slide, 4.55, 0.34, 8.40, 1.10,
                [(title, 19, True, BLUE),
                 (headline, 12.5, False, INK2, 2),
                 (takeaway, 11.5, False, MUTED, 3)])
        add_fitted_picture(slide, FIG / image, 0.45, img_top, 12.44, img_h)
    else:
        add_box(slide, 4.55, 0.44, 8.30, 0.92,
                [(title, 19, True, BLUE), (headline, 12.5, False, INK2, 2)])
        add_fitted_picture(slide, FIG / image, 0.45, img_top, 12.44, img_h)
        add_box(slide, 0.62, img_top + img_h + 0.06, 12.10, 0.58,
                [(takeaway, 12.5, False, INK2)])


# --------------------------------------------------------------------------- build
prs = Presentation(str(TEMPLATE))

# duplicate the results slide (index 6) four times -> 9 results slides in total
for _ in range(4):
    duplicate_slide(prs, 6)
reorder(prs, [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 14, 15, 16, 17, 10, 11, 12, 13])
S = list(prs.slides)

# ---------------------------------------------------------------- 1 title
s = S[0]
title = shape_by_text(s, 'Project Title')
title.top, title.height = Inches(2.58), Inches(1.24)      # 30pt over two lines needs the room
set_text(title, [('Seasonal Agriculture', 30, True, INK),
                 ('Performance Analysis', 30, True, TEAL)])
name_box = shape_by_text(s, 'Student Name')
name_box.left, name_box.width = Inches(0.78), Inches(5.20)   # template had it half off-slide
set_text(name_box,
         [('Shivani Chaudhary', 15, True, INK),
          ('Birla Institute of Technology and Science, Pilani – K. K. Birla Goa Campus', 11, False, INK2)])
aicte = shape_by_text(s, 'AICTE STU ID')
aicte.left, aicte.top, aicte.width = Inches(0.78), Inches(4.92), Inches(5.20)
set_text(aicte, [('AICTE STU ID : STU6a6a3b03e36e41785346819', 11.5, False, INK2)])
for shp in s.shapes:
    if shp.has_text_frame and shp.text_frame.text.strip() == '' and shp.name == 'Text Placeholder 1':
        set_text(shp, [('VOIS AICTE Batch 1  ·  2026–2027', 12, True, BLUE),
                       ('Major Project  ·  Course: Data Visualization', 11, False, MUTED, 2),
                       ('4,000 farms  ·  8 crops  ·  8 states  ·  3 seasons', 11, False, MUTED, 2)])

# ---------------------------------------------------------------- 2 problem statement
s = S[1]
body = shape_by_text(s, 'Agricultural activities are influenced')
set_text(body, [
    ('Agricultural activities are influenced by seasonal variations in environmental '
     'conditions, farming practices, resource availability and market conditions. '
     'As a result, agricultural performance may differ from one season to another.',
     15, False, INK2),
    ('However, raw agricultural data does not clearly explain how agricultural performance '
     'changes across seasons or what patterns can be observed in different seasonal conditions.',
     15, False, INK2, 10),
    ('The problem is to analyse the given agricultural dataset and investigate seasonal '
     'differences in agricultural performance by identifying meaningful patterns, trends, '
     'relationships and variations within the available data.',
     15, True, INK, 10),
])

# ---------------------------------------------------------------- 3 project description
s = S[2]
set_text(shape_by_text(s, 'Project Description'),
         [('Project Description', 30, True, INK)])
add_box(s, 0.72, 1.72, 11.90, 4.90, [
    ('The project analyses 4,000 Indian farm records across the three cropping seasons — '
     'Kharif (monsoon), Rabi (winter) and Zaid (summer) — spanning 8 crops, 8 states and '
     '28 variables covering weather, soil, farming practice, output, water use and economics.',
     14, False, INK2),

    ('Approach', 16, True, BLUE, 14),
    ('1.  Data-quality audit — verified four arithmetic identities in the data, which allowed '
     '32 missing yields to be reconstructed exactly rather than estimated.', 13, False, INK2, 6),
    ('2.  Crop-relative yield — the key engineered feature. Sugarcane yields 48 t/ha and pulses '
     '0.9 t/ha, so raw averages measure crop mix, not seasonal performance. Dividing each farm '
     'by its own crop median puts all 8 crops on one comparable scale.', 13, False, INK2, 6),
    ('3.  Hypothesis testing — Kruskal-Wallis, pairwise Mann-Whitney with Bonferroni correction, '
     'and chi-square, with effect sizes reported alongside every p-value.', 13, False, INK2, 6),
    ('4.  Robustness — the seasonal ranking is re-tested inside all 30 sub-groups '
     '(crop, state, district, irrigation method) to rule out a composition artefact.', 13, False, INK2, 6),
    ('5.  Modelling — nested OLS models separate what the season label explains from what '
     'measurable growing conditions explain.', 13, False, INK2, 6),

    ('Headline result', 16, True, BLUE, 14),
    ('Productivity falls 30 % from Kharif to Zaid and the ranking holds in 30 of 30 sub-groups. '
     'The cause is a rainfall optimum near 760 mm, not the calendar. Because cost per hectare is '
     'statistically identical in every season while revenue falls, loss-making farms rise from '
     '42 % in Kharif to 65 % in Zaid.', 13, True, INK, 6),
])

# ---------------------------------------------------------------- 4 end users
s = S[3]
set_text(shape_by_text(s, 'END USERS'), [('WHO ARE THE END USERS?', 30, True, INK)])
body = [shp for shp in s.shapes
        if shp.is_placeholder and shp.placeholder_format.idx == 12][0]
body.left, body.top   = Inches(0.72), Inches(1.90)
body.width, body.height = Inches(11.90), Inches(4.90)
set_text(body, [
    ('Farmers and Farmer Producer Organisations (FPOs)', 15, True, TEAL),
    ('Which crop to sow in which season, and when the expected return will not cover the input cost.',
     12.5, False, INK2, 2),

    ('Agricultural extension officers and Krishi Vigyan Kendras', 15, True, BLUE, 10),
    ('Season-specific advisories — irrigation targets, nutrient balance and Kharif disease control.',
     12.5, False, INK2, 2),

    ('State agriculture departments and policy planners', 15, True, ORANGE, 10),
    ('Where to direct irrigation subsidies, drainage investment and crop-diversification support.',
     12.5, False, INK2, 2),

    ('Agri-input and irrigation companies', 15, True, TEAL, 10),
    ('Evidence that drip systems hold up best in the water-scarce Zaid season.',
     12.5, False, INK2, 2),

    ('Agricultural insurers and rural lenders', 15, True, BLUE, 10),
    ('Season- and crop-level loss rates (42 % → 65 %) for pricing risk and structuring credit.',
     12.5, False, INK2, 2),

    ('Agri-tech platforms and researchers', 15, True, ORANGE, 10),
    ('A reproducible baseline for seasonal benchmarking and yield-advisory products.',
     12.5, False, INK2, 2),
])

# ---------------------------------------------------------------- 5 technology used
s = S[4]
set_text(shape_by_text(s, 'Technology Used'), [('Technology Used', 30, True, INK)])
body = [shp for shp in s.shapes
        if shp.is_placeholder and shp.placeholder_format.idx == 12][0]
body._element.getparent().remove(body._element)

add_box(s, 2.10, 1.55, 5.30, 5.30, [
    ('Language & environment', 16, True, BLUE),
    ('Python 3  ·  Jupyter Notebook', 13, False, INK2, 4),

    ('Data handling', 16, True, BLUE, 12),
    ('pandas — cleaning, grouping, pivot tables', 13, False, INK2, 4),
    ('NumPy — vectorised computation, least-squares', 13, False, INK2, 2),

    ('Statistics', 16, True, BLUE, 12),
    ('SciPy — Kruskal-Wallis, Mann-Whitney U,', 13, False, INK2, 4),
    ('ANOVA, chi-square, Spearman correlation', 13, False, INK2, 0),
    ('OLS regression implemented with NumPy', 13, False, INK2, 2),

    ('Visualisation', 16, True, BLUE, 12),
    ('Matplotlib and Seaborn — a colour-vision-safe', 13, False, INK2, 4),
    ('palette validated to ΔE ≥ 9 separation', 13, False, INK2, 0),
])
add_box(s, 7.75, 1.55, 5.10, 5.30, [
    ('Analytical techniques applied', 16, True, TEAL),
    ('•  Data-quality auditing via arithmetic identity checks', 12.5, False, INK2, 5),
    ('•  Group-wise median imputation', 12.5, False, INK2, 3),
    ('•  Normalisation to remove crop-scale confounding', 12.5, False, INK2, 3),
    ('•  Non-parametric hypothesis testing', 12.5, False, INK2, 3),
    ('•  Multiple-comparison (Bonferroni) correction', 12.5, False, INK2, 3),
    ('•  Effect sizes: η², Cramér\'s V, P(A > B)', 12.5, False, INK2, 3),
    ('•  Sub-group robustness testing (30 groups)', 12.5, False, INK2, 3),
    ('•  Partial correlation to expose a confound', 12.5, False, INK2, 3),
    ('•  Polynomial response-curve fitting', 12.5, False, INK2, 3),
    ('•  Nested regression models for mediation', 12.5, False, INK2, 3),
    ('•  Break-even / unit-economics analysis', 12.5, False, INK2, 3),
])

# ---------------------------------------------------------------- 6-14 results
R = [
    dict(title='Seasonal performance at a glance',
         headline='Six measures, three seasons — nature varies enormously, farmer spending does not',
         image='fig09_dashboard.png',
         takeaway='Rainfall differs 3× between seasons and yield falls 30 %, yet cost per hectare '
                  'varies by under 1 %. That mismatch is where the losses come from.'),
    dict(title='Data quality, cleaning and the key feature',
         headline='Four arithmetic identities verified — so 32 missing yields were reconstructed, not guessed',
         image='code01_cleaning.png',
         takeaway='Crop-relative yield collapses a 51× range between sugarcane and pulses onto a '
                  'common 1.0 scale, so seasons compare fairly.',
         img_top=1.46, img_h=5.52, takeaway_on_top=True),
    dict(title='Productivity across seasons',
         headline='Kharif 1.13  ▸  Rabi 0.97  ▸  Zaid 0.79 — a 30 % gap from monsoon to summer',
         image='fig02_yield_by_season.png',
         takeaway='Kruskal-Wallis H = 254, p < 1e-55. All three pairwise gaps survive Bonferroni '
                  'correction; a Kharif farm out-yields a Zaid farm 70 % of the time.'),
    dict(title='Is the ranking real, or a crop-mix artefact?',
         headline='Kharif > Rabi > Zaid holds in 8/8 crops, 8/8 states, 10/10 districts, 4/4 irrigation methods',
         image='fig03_robustness_by_crop.png',
         takeaway='30 out of 30 independent sub-groups agree. The penalty size varies (pulses −37 %, '
                  'wheat −22 %) but its direction never reverses — this is not Simpson\'s paradox.'),
    dict(title='Why the gap exists',
         headline='Water availability — not the calendar — is what actually moves yield',
         image='fig04_rainfall_response.png',
         takeaway='Yield peaks near 760 mm. 62 % of Kharif farms sit above the optimum, so extra rain '
                  'costs them yield; no Zaid farm ever reaches it. Water supply — not the calendar — is the mechanism.'),
    dict(title='Resource usage and water efficiency',
         headline='Identical inputs applied in every season, but 30 % less output per cubic metre',
         image='fig05_water_efficiency.png',
         takeaway='Fertiliser, N-P-K, pesticide, seed quality and water applied are statistically '
                  'indistinguishable across seasons. Input behaviour is season-blind; output is not.'),
    dict(title='Economic outcomes',
         headline='Farmers spend the same in every season; only what they earn back changes',
         image='fig06_economics.png',
         takeaway='Cost/ha (₹66.3k–₹66.7k, p = 0.07) and market price (p = 0.94) are flat. '
                  'Loss-making farms rise from 42 % in Kharif to 65 % in Zaid (χ² = 92.7, p < 1e-20).'),
    dict(title='Break-even by crop and season',
         headline='Wheat, rice and maize sit below break-even in all three seasons',
         image='fig07_breakeven_heatmap.png',
         takeaway='Sugarcane and chilli clear break-even by 2.3–2.4× in every season. Wheat, rice and '
                  'maize sit below break-even in all three — only 8 % of Zaid wheat farms make money.'),
    dict(title='Modelling the drivers',
         headline='Growing conditions explain twice what the season label does — season is a proxy',
         image='code02_testing.png',
         takeaway='41 % of the Zaid penalty is mediated by measurable conditions — overwhelmingly '
                  'rainfall — and conditions can be managed where a season cannot.',
         img_top=1.46, img_h=5.52, takeaway_on_top=True),
]
for slide, cfg in zip(S[5:14], R):
    results_slide(slide, **cfg)

# ---------------------------------------------------------------- 15 future scope
s = S[14]
set_text(shape_by_text(s, 'Future scope'), [('Future Scope', 30, True, INK)])
for shp in list(s.shapes):
    if shp.has_text_frame and shp.text_frame.text.strip() == '' and shp.name.startswith('Text Placeholder 30'):
        shp._element.getparent().remove(shp._element)
add_box(s, 0.72, 1.45, 5.90, 5.40, [
    ('Extending the analysis', 16, True, BLUE),
    ('Add a time dimension.  This dataset has no year or date, so no trend or forecast is '
     'possible. Multi-year data would turn seasonal differences into seasonal trends.',
     12.5, False, INK2, 6),
    ('Predictive modelling.  Random Forest / gradient boosting on yield and profit, with '
     'SHAP values to confirm the rainfall optimum non-parametrically.', 12.5, False, INK2, 6),
    ('Real geography.  Join to genuine district coordinates and IMD rainfall grids to map '
     'where the 760 mm optimum is actually met.', 12.5, False, INK2, 6),
    ('Farm clustering.  Segment farms by input strategy to find which combinations survive Zaid.',
     12.5, False, INK2, 6),
])
add_box(s, 6.95, 1.45, 5.70, 5.40, [
    ('Turning it into a product', 16, True, TEAL),
    ('Season-aware advisory tool.  Enter crop, district and irrigation method; return the '
     'expected break-even yield and a season-specific input budget.', 12.5, False, INK2, 6),
    ('Irrigation-gap dashboard.  Flag farms below the 760 mm water target and quantify the '
     'yield recoverable by closing that gap.', 12.5, False, INK2, 6),
    ('Loss early-warning.  A classifier for loss-making crop-season combinations, usable by '
     'lenders and insurers for risk pricing.', 12.5, False, INK2, 6),
    ('Live data integration.  Connect IMD weather feeds and mandi price APIs to move from '
     'retrospective analysis to in-season decision support.', 12.5, False, INK2, 6),
])

# ---------------------------------------------------------------- 16 github
s = S[15]
set_text(shape_by_text(s, 'GitHub Link'), [('GitHub Link', 30, True, INK)])
link = shape_by_text(s, 'github.com')
set_text(link, [('https://github.com/shivanichaudhary2405/Seasonal-Agriculture-Performance-Analysis',
                 15, True, BLUE)])
link.width, link.top = Inches(11.9), Inches(1.75)
for shp in list(s.shapes):
    if shp.has_text_frame and shp.text_frame.text.strip() == '' and shp.name.startswith('Text Placeholder 30'):
        shp._element.getparent().remove(shp._element)
add_box(s, 0.80, 2.60, 11.90, 4.10, [
    ('Repository contents', 16, True, TEAL),
    ('Seasonal_Agriculture_Performance_Analysis.ipynb — the complete analysis, executed with all outputs',
     13, False, INK2, 6),
    ('data/  — the source dataset (unmodified)', 13, False, INK2, 3),
    ('figures/  — the 9 charts and 2 code screenshots used in this deck', 13, False, INK2, 3),
    ('outputs/  — cleaned dataset and 5 result tables (CSV)', 13, False, INK2, 3),
    ('README.md  ·  requirements.txt', 13, False, INK2, 3),
    ('⚠  Replace the link above with your own repository URL before submitting.',
     12.5, True, RED, 14),
])

# ---------------------------------------------------------------- 17 certificate
s = S[16]
set_text(shape_by_text(s, 'certificate'),
         [('VOIS Course Completion Certificate', 26, True, INK),
          ('Course name - Data Visualization', 17, False, INK2, 4)])
for shp in list(s.shapes):
    if shp.has_text_frame and shp.text_frame.text.strip() == '' and shp.name.startswith('Text Placeholder 30'):
        shp._element.getparent().remove(shp._element)
add_box(s, 0.80, 2.30, 11.90, 3.60, [
    ('⚠  Paste your VOIS "Data Visualization" course completion certificate image on this slide.',
     15, True, RED),
    ('Download it from your VOIS learner dashboard, then Insert ▸ Picture and centre it here.',
     13, False, INK2, 8),
], align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 18 thank you
s = S[17]
set_text(shape_by_text(s, 'Thank you'), [('Thank you', 40, True, INK)])
for shp in list(s.shapes):
    if shp.has_text_frame and shp.text_frame.text.strip() in {'.', ''}:
        shp._element.getparent().remove(shp._element)
add_box(s, 0.80, 2.70, 11.75, 2.40, [
    ('Nature varies enormously between seasons; farmer spending does not.',
     21, True, INK),
    ('That mismatch — not the weather alone — is what turns 42 % of Kharif farms '
     'into 65 % of Zaid farms making a loss.', 14, False, INK2, 12),
    ('Shivani Chaudhary   ·   BITS Pilani, K. K. Birla Goa Campus   ·   STU6a6a3b03e36e41785346819', 12, False, MUTED, 18),
], align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print(f'{len(S)} slides -> {OUT}')
